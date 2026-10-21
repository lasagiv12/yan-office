#!/usr/bin/env python3
"""Client dashboard: the Excel dashboard the consultant works in, and on request a client-ready PDF from it.

    python dashboard.py build data.json --xlsx "דשבורד 2026.xlsx"        create or update the Excel dashboard
    python dashboard.py read "דשבורד 2026.xlsx" [--json data.json]       the Excel back as a data file
    python dashboard.py report "דשבורד 2026.xlsx" --pdf out.pdf [--html out.html]
                                                                          the presentation report, from the Excel as it is now

The Excel file is the source of truth. To add or fix a month: `read`, change the data, `build` again (the whole
file is written from scratch, because openpyxl drops charts when it opens an existing file). Everything that
differs between businesses lives in the data file: labels, which expenses to show, what activity is measured.
Calculated values are stored next to every formula, so the file shows numbers even where nothing recalculates
(previews, Protected View). Needs only openpyxl. Standalone on purpose: it ships inside the plugin.

Data file (JSON):
  client, consultant, year, updated (YYYY-MM-DD)
  labels: revenue, direct_costs (null = the business has none), expenses, profit, owner_salary (null = none)
  expense_groups: up to 4 expense lines to show apart from the owner's salary, e.g. ["פרסום ממומן"]
  activity: {title, items: [{label, unit}], total: label or null}  up to 6 items; total only when adding them up
            means something (treatments + courses: yes; customers + orders: no)
  months: [{month: "YYYY-MM", revenue, direct_costs, expenses, owner_salary, groups: [...], activity: [...],
            source}]  amounts without VAT; an activity value may be null when it was not measured that month
  insights: [text], notes: [[title, text]]
"""
import argparse
import base64
import datetime as dt
import html
import json
import math
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape as xml_escape

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.chart.data_source import AxDataSource, NumData, NumVal, StrData, StrRef, StrVal
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
RTL = True  # charts read right to left, like the sheets: first month on the right
MONTHS_HE = ["ינואר", "פברואר", "מרץ", "אפריל", "מאי", "יוני", "יולי", "אוגוסט", "ספטמבר", "אוקטובר", "נובמבר", "דצמבר"]
SHORT_HE = ["ינו׳", "פבר׳", "מרץ", "אפר׳", "מאי", "יוני", "יולי", "אוג׳", "ספט׳", "אוק׳", "נוב׳", "דצמ׳"]
DEFAULT_LABELS = {"revenue": "הכנסות", "direct_costs": "עלויות ישירות", "expenses": "הוצאות חודשיות",
                  "profit": "רווח לפני מס", "owner_salary": "שכר בעלים"}

# Palette: categorical slots in fixed order (validated set), ink and surfaces.
SERIES = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "6250D6"]
INK, INK2, MUTED = "17233B", "52514E", "898781"
GRID, BASE = "E1E0D9", "C3C2B7"
GOOD, BAD = "006300", "C62828"
ACCENT = "1F3A5F"

# ---------------------------------------------------------------- data

def month_label(key, with_year=False):
    y, m = key.split("-")
    return f"{MONTHS_HE[int(m) - 1]} {y}" if with_year else MONTHS_HE[int(m) - 1]


def months_count(n):
    return {1: "חודש אחד", 2: "חודשיים"}.get(n, f"{n} חודשים")


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def normalize(data):
    """Fill defaults and check the data. Exits with every problem listed, in Hebrew."""
    labels = {**DEFAULT_LABELS, **(data.get("labels") or {})}
    for k in ("direct_costs", "owner_salary"):
        if not labels.get(k):
            labels[k] = None
    data["labels"] = labels
    data.setdefault("expense_groups", [])
    act = data.get("activity") or {}
    if "activities" in data and not act:  # older demo files
        act = {"title": "פעילות", "items": [{"label": a, "unit": ""} for a in data.pop("activities")], "total": "סה״כ פעילות"}
    act = {"title": act.get("title") or "פעילות", "items": [i if isinstance(i, dict) else {"label": i, "unit": ""}
                                                            for i in act.get("items", [])], "total": act.get("total") or None}
    data["activity"] = act
    data.setdefault("insights", [])
    data.setdefault("notes", [])
    groups, items = data["expense_groups"], act["items"]
    errors = []
    if len(groups) > 4:
        errors.append("אפשר עד 4 סעיפי הוצאה להצגה בנפרד")
    if len(items) > 6:
        errors.append("אפשר עד 6 מדדי פעילות")
    if not data.get("client"):
        errors.append("חסר שם הלקוח")
    seen = set()
    for m in data.get("months", []):
        k = m.get("month", "?")
        if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", str(k)):
            errors.append(f"{k}: החודש צריך להיות בפורמט YYYY-MM")
            continue
        if k in seen:
            errors.append(f"{k}: החודש מופיע פעמיים")
        seen.add(k)
        m.setdefault("direct_costs", 0)
        m.setdefault("owner_salary", 0)
        m.setdefault("groups", [])
        m.setdefault("activity", [])
        m.setdefault("source", "")
        for f in ("revenue", "expenses", "direct_costs", "owner_salary"):
            if num(m.get(f)) is None:
                errors.append(f"{k}: חסר ערך מספרי ({labels.get(f) or f})")
        if not labels["direct_costs"]:
            m["direct_costs"] = 0
        if not labels["owner_salary"]:
            m["owner_salary"] = 0
        if len(m["groups"]) != len(groups):
            errors.append(f"{k}: {len(m['groups'])} סעיפי הוצאה במקום {len(groups)}")
        elif any(num(v) is None for v in m["groups"]):
            errors.append(f"{k}: סעיף הוצאה בלי ערך מספרי")
        if len(m["activity"]) != len(items):
            errors.append(f"{k}: {len(m['activity'])} מדדי פעילות במקום {len(items)}")
        if not errors and m["owner_salary"] + sum(m["groups"]) > m["expenses"] + 0.5:
            errors.append(f"{k}: {labels['owner_salary'] or 'שכר'} והסעיפים גדולים מסך ההוצאות")
    if not data.get("months"):
        errors.append("אין אף חודש בנתונים")
    if errors:
        sys.exit("שגיאות בנתונים:\n- " + "\n- ".join(errors))
    data["months"].sort(key=lambda m: m["month"])
    return data


def derive(data):
    """Everything the views show, computed once. The Excel computes the same with formulas."""
    ms, items = data["months"], data["activity"]["items"]
    rows = []
    for i, m in enumerate(ms):
        profit = m["revenue"] - m["direct_costs"] - m["expenses"]
        prev = rows[i - 1] if i else None
        acts = [num(v) for v in m["activity"]]
        rows.append({
            **m, "activity": acts,
            "label": month_label(m["month"]),
            "profit": profit,
            "margin": profit / m["revenue"] if m["revenue"] else 0,
            "rest": m["expenses"] - m["owner_salary"] - sum(m["groups"]),
            "total_activity": sum(acts) if acts and None not in acts else None,
            "d_revenue": None if prev is None else m["revenue"] - prev["revenue"],
            "d_profit": None if prev is None else profit - prev["profit"],
        })
    n = len(rows)
    avg = lambda key: sum(r[key] for r in rows) / n

    def avg_act(j):
        vals = [r["activity"][j] for r in rows if r["activity"][j] is not None]
        return sum(vals) / len(vals) if vals else None

    pick = lambda key, f: f(rows, key=lambda r: r[key])["label"]
    return {
        "rows": rows, "n": n,
        "avg": {k: avg(k) for k in ("revenue", "expenses", "profit", "owner_salary", "direct_costs")},
        "avg_activity": [avg_act(j) for j in range(len(items))],
        "forecast": {"revenue": avg("revenue") * 12, "expenses": (avg("expenses") + avg("direct_costs")) * 12,
                     "profit": avg("profit") * 12},
        "best_revenue": pick("revenue", max), "best_profit": pick("profit", max), "worst_profit": pick("profit", min),
        "period": (month_label(rows[0]["month"], True) if n == 1 else
                   f"{month_label(rows[0]['month'])} – {month_label(rows[-1]['month'], True)}"),
    }


def expense_parts(data):
    """(name, getter) for every expense part shown apart: owner's salary, the chosen lines, everything else."""
    L = data["labels"]
    parts = [(L["owner_salary"], lambda r: r["owner_salary"])] if L["owner_salary"] else []
    parts += [(g, (lambda j: lambda r: r["groups"][j])(j)) for j, g in enumerate(data["expense_groups"])]
    if parts:
        parts.append(("שאר ההוצאות", lambda r: r["rest"]))
    return parts

# ---------------------------------------------------------------- Excel

FONT = "Arial"
thin = Side(style="thin", color=GRID)
fill = lambda c: PatternFill("solid", start_color=c, end_color=c)
INPUT_FILL, HEAD_FILL, CARD_FILL = fill("FFF8E1"), fill(ACCENT), fill("F4F6FA")
MONEY = '#,##0" ₪";[Red]-#,##0" ₪";0" ₪"'
DELTA = '"▲ "#,##0" ₪";"▼ "#,##0" ₪";0" ₪"'
PCT = '0.0%;[Red]-0.0%'
COUNT = '#,##0'
DATA, SETTINGS, INSIGHTS, DASH = "חודש בחודש", "הגדרות", "תובנות", "דשבורד"  # Numbers can't follow references to some Hebrew names, e.g. "נתונים"
HDR, FIRST = 4, 5  # header row and first month row in the data sheet


class Cache:
    """Calculated values for formula cells and chart series, written into the file after openpyxl saves it."""

    def __init__(self):
        self.cells = {}

    def put(self, ws, coord, formula, value):
        ws[coord] = formula
        self.cells[(ws.title, coord)] = value
        return ws[coord]


def avg_format(v):
    return "#,##0.0" if v is None or abs(v) < 100 else "#,##0"


def data_columns(data):
    L, items = data["labels"], data["activity"]["items"]
    cols = [("month_label", "חודש", "text"), ("revenue", f"{L['revenue']} ללא מע״מ", "in")]
    if L["direct_costs"]:
        cols.append(("direct_costs", L["direct_costs"], "in"))
    cols += [("expenses", f"{L['expenses']} ללא מע״מ", "in"), ("profit", L["profit"], "f"), ("margin", "שיעור רווח", "f")]
    if L["owner_salary"]:
        cols.append(("owner_salary", L["owner_salary"], "in"))
    cols += [(f"g{j}", g, "in") for j, g in enumerate(data["expense_groups"])]
    if L["owner_salary"] or data["expense_groups"]:
        cols.append(("rest", "שאר ההוצאות", "f"))
    cols += [(f"a{j}", it["label"], "in") for j, it in enumerate(items)]
    if items and data["activity"]["total"]:
        cols.append(("total_activity", data["activity"]["total"], "f"))
    cols += [("source", "קובץ מקור", "text"), ("month", "מזהה חודש", "text")]
    return [(key, title, kind, get_column_letter(i + 1)) for i, (key, title, kind) in enumerate(cols)]


def row_value(r, key):
    if key.startswith("g") and key[1:].isdigit():
        return r["groups"][int(key[1:])]
    if key.startswith("a") and key[1:].isdigit():
        return r["activity"][int(key[1:])]
    if key == "month_label":
        return r["label"]
    return r.get(key)


def build_data_sheet(wb, data, d, cache):
    ws = wb.create_sheet(DATA)
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    cols = data_columns(data)
    col = {k: letter for k, _, _, letter in cols}
    ws["A1"] = "חודש בחודש: שורה לכל חודש"
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color=INK)
    ws["A2"] = "תאים בצבע שמנת מזינים מהניתוח הפיננסי. השאר מחושב לבד. כל שינוי כאן מתעדכן בדשבורד."
    ws["A2"].font = Font(name=FONT, size=10, color=INK2)
    for key, title, kind, letter in cols:
        c = ws[f"{letter}{HDR}"]
        c.value = title
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[letter].width = 34 if key == "source" else 14
    ws.row_dimensions[HDR].height = 36
    items = data["activity"]["items"]
    for i, r in enumerate(d["rows"]):
        row = FIRST + i
        for key, _, kind, letter in cols:
            coord = f"{letter}{row}"
            if key == "profit":
                direct = f"-{col['direct_costs']}{row}" if "direct_costs" in col else ""
                c = cache.put(ws, coord, f"={col['revenue']}{row}{direct}-{col['expenses']}{row}", r["profit"])
            elif key == "margin":
                c = cache.put(ws, coord, f"=IF({col['revenue']}{row}=0,0,{col['profit']}{row}/{col['revenue']}{row})", r["margin"])
            elif key == "rest":
                parts = (["owner_salary"] if "owner_salary" in col else []) + [f"g{j}" for j in range(len(data["expense_groups"]))]
                c = cache.put(ws, coord, f"={col['expenses']}{row}-" + "-".join(f"{col[p]}{row}" for p in parts), r["rest"])
            elif key == "total_activity":
                a0, a1 = col["a0"], col[f"a{len(items) - 1}"]
                c = cache.put(ws, coord, f'=IF(COUNT({a0}{row}:{a1}{row})={len(items)},SUM({a0}{row}:{a1}{row}),"")',
                              "" if r["total_activity"] is None else r["total_activity"])
            else:
                c = ws[coord]
                c.value = row_value(r, key)
            c.font = Font(name=FONT, size=11, color=INK if kind != "f" else INK2, bold=key == "month_label")
            c.border = Border(bottom=thin)
            c.alignment = Alignment(horizontal="right" if key == "source" else "center")
            if kind == "in":
                c.fill = INPUT_FILL
            if key == "margin":
                c.number_format = PCT
            elif key.startswith("a") and key[1:].isdigit() or key == "total_activity":
                c.number_format = COUNT
            elif kind in ("in", "f"):
                c.number_format = MONEY
    ws.freeze_panes = f"B{FIRST}"
    end = FIRST + d["n"] - 1
    L = data["labels"]
    note = f"{L['profit']} = {L['revenue']}" + (f" − {L['direct_costs']}" if L["direct_costs"] else "") + f" − {L['expenses']}."
    ws[f"A{end + 2}"] = note
    ws[f"A{end + 2}"].font = Font(name=FONT, size=9, color=MUTED)
    ws.column_dimensions["A"].width = 12
    return col, end


def build_settings_sheet(wb, data):
    ws = wb.create_sheet(SETTINGS)
    ws.sheet_view.rightToLeft = True
    config = {k: data[k] for k in ("client", "consultant", "year", "updated", "labels", "expense_groups", "activity", "notes")
              if k in data}
    ws["A1"] = "הגדרות הדשבורד של הלקוח. העוזר קורא אותן בכל עדכון. לא לשנות ידנית."
    ws["A2"] = "config"
    ws["B2"] = json.dumps(config, ensure_ascii=False)
    ws.sheet_state = "hidden"


def build_insights_sheet(wb, data):
    ws = wb.create_sheet(INSIGHTS)
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws["A1"] = "תובנות לפעולה"
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color=INK)
    ws["A2"] = "אפשר לערוך כאן את הטקסט. הדשבורד מציג אותו כפי שהוא."
    ws["A2"].font = Font(name=FONT, size=10, color=INK2)
    for i, text in enumerate(data["insights"]):
        ws.cell(4 + i, 1, i + 1).font = Font(name=FONT, size=11, bold=True, color=ACCENT)
        c = ws.cell(4 + i, 2, text)
        c.font = Font(name=FONT, size=11, color=INK)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[4 + i].height = 32
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 110


def _num_cache(values):
    return NumData(ptCount=len(values), pt=[NumVal(idx=i, v=v) for i, v in enumerate(values) if v not in (None, "")])


def _str_ref(f, values):
    return AxDataSource(strRef=StrRef(f=f, strCache=StrData(ptCount=len(values), pt=[StrVal(idx=i, v=v) for i, v in enumerate(values)])))


def _series(ws, col, r0, r1, title, values, cats_f, cats, color, line=False):
    s = Series(Reference(ws, min_col=col, min_row=r0, max_row=r1), title=title)
    s.val.numRef.numCache = _num_cache(values)
    s.cat = _str_ref(cats_f, cats)
    if line:
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 28575
        s.marker.symbol = "circle"
        s.marker.size = 7
        s.marker.graphicalProperties = GraphicalProperties(solidFill=color)
        s.marker.graphicalProperties.line.solidFill = "FFFFFF"
        s.smooth = False
    else:
        s.graphicalProperties.solidFill = color
        s.graphicalProperties.line.noFill = True
    return s


def _chart_frame(ch, title, height=8.2, fmt="#,##0"):
    ch.title = title
    ch.title_text = title
    ch.graphical_properties = GraphicalProperties(solidFill="FFFFFF")  # opaque, so the picture underneath stays hidden
    ch.legend.position = "b"
    ch.height, ch.width = height, 16.5
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.numFmt = fmt
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRID))
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=BASE))
    ch.x_axis.tickLblPos = "low"
    if RTL:
        ch.x_axis.scaling.orientation = "maxMin"  # months from right to left, value axis on the right
    return ch


def build_dashboard_sheet(wb, data, d, col, end, cache):
    ws = wb.create_sheet(DASH, 0)
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.column_dimensions["A"].width = 2
    for i in range(2, 18):
        ws.column_dimensions[get_column_letter(i)].width = 11.5
    L, act, groups, rows, n = data["labels"], data["activity"], data["expense_groups"], d["rows"], d["n"]
    items = act["items"]
    dsheet = wb[DATA]
    rng = lambda key: f"'{DATA}'!${col[key]}${FIRST}:${col[key]}${end}"
    dref = lambda key, row: f"'{DATA}'!{col[key]}{row}"

    def text(cell, value, size=11, bold=False, color=INK, align="right", wrap=False):
        c = ws[cell]
        c.value = value
        c.font = Font(name=FONT, size=size, bold=bold, color=color)
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        return c

    def style(c, size=11, bold=False, color=INK, align="center", fmt=None):
        c.font = Font(name=FONT, size=size, bold=bold, color=color)
        c.alignment = Alignment(horizontal=align, vertical="center")
        if fmt:
            c.number_format = fmt
        return c

    def section(row, title, sub=None):
        text(f"B{row}", title, 13, True, ACCENT)
        if sub:
            ws.merge_cells(f"H{row}:Q{row}")
            text(f"H{row}", sub, 9, color=MUTED, align="left")
        for i in range(2, 18):
            ws.cell(row, i).border = Border(bottom=Side(style="medium", color=ACCENT))

    def cards(row, specs, per_row=4):
        """specs: (label, formula, value, number_format, caption). Returns the next free row."""
        width = 16 // per_row
        for i, (label, formula, value, fmt, caption) in enumerate(specs):
            r = row + (i // per_row) * 4
            c0 = 2 + (i % per_row) * width
            c1 = c0 + width - 2  # last merged column; one gutter column after it
            a, b = get_column_letter(c0), get_column_letter(c1)
            for rr in range(r, r + 3):
                ws.merge_cells(f"{a}{rr}:{b}{rr}")
                for cc in range(c0, c1 + 1):
                    ws.cell(rr, cc).fill = CARD_FILL
            for cc in range(c0, c1 + 1):
                ws.cell(r, cc).border = Border(top=Side(style="thick", color=ACCENT if i == 2 and per_row != 3 else SERIES[0]))
            text(f"{a}{r}", label, 10, True, INK2)
            v = cache.put(ws, f"{a}{r + 1}", formula, value)
            style(v, 20, True, align="right", fmt=fmt)
            text(f"{a}{r + 2}", caption, 9, color=MUTED)
            ws.row_dimensions[r + 1].height = 30
        return row + ((len(specs) - 1) // per_row + 1) * 4

    def head(row, c, title):
        x = ws.cell(row, c, title)
        x.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        x.fill = HEAD_FILL
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 30

    def table(row, title, columns):
        """columns: (header, key or callable(i) -> (formula, value), number_format). Returns (header row, next free row)."""
        text(f"B{row}", title, 11, True, INK)
        h = row + 1
        for j, (header, _, _) in enumerate(columns):
            head(h, 2 + j, header)
        for i in range(n):
            src, r = FIRST + i, h + 1 + i
            for j, (_, key, fmt) in enumerate(columns):
                coord = f"{get_column_letter(2 + j)}{r}"
                if callable(key):
                    formula, value = key(i)
                else:
                    formula, value = f"={dref(key, src)}", row_value(rows[i], key)
                c = cache.put(ws, coord, formula, "" if value is None else value) if str(formula).startswith("=") else ws[coord]
                if not str(formula).startswith("="):
                    c.value = formula
                style(c, bold=j == 0 or key == "total_activity", fmt=fmt)
                c.border = Border(bottom=thin)
        return h, h + n + 2

    # Header
    ws.merge_cells("B2:Q2")
    text("B2", f"דשבורד ניהולי {data.get('year', '')} | {data['client']}", 22, True, INK)
    ws.row_dimensions[2].height = 34
    ws.merge_cells("B3:Q3")
    updated = dt.date.fromisoformat(data["updated"]).strftime("%d/%m/%Y") if data.get("updated") else ""
    text("B3", f"{d['period']} · {months_count(n)} · עודכן {updated} · הוכן על ידי {data.get('consultant', '')}", 11, color=INK2)

    av = d["avg"]
    section(5, "התמונה הכספית", f"ממוצע לחודש, לפי {months_count(n)}")
    money = [(L["revenue"], f"=AVERAGE({rng('revenue')})", av["revenue"], MONEY, "ללא מע״מ"),
             (L["expenses"], f"=AVERAGE({rng('expenses')})", av["expenses"], MONEY, "ללא מע״מ"),
             (L["profit"], f"=AVERAGE({rng('profit')})", av["profit"], MONEY, "אחרי כל ההוצאות")]
    if L["owner_salary"]:
        money.append((L["owner_salary"], f"=AVERAGE({rng('owner_salary')})", av["owner_salary"], MONEY, "כלול בהוצאות"))
    row = cards(6, money, per_row=len(money)) + 1
    if items:
        section(row, act["title"], f"ממוצע לחודש, לפי {months_count(n)}")
        specs = [(it["label"], f'=IFERROR(AVERAGE({rng(f"a{j}")}),"")', "" if d["avg_activity"][j] is None else d["avg_activity"][j],
                  avg_format(d["avg_activity"][j]), f"{it.get('unit') or ''} בחודש".strip()) for j, it in enumerate(items)]
        row = cards(row + 1, specs, per_row=4 if len(items) != 3 else 3) + 1

    # Tables
    section(row, "לוח שינוי חודשי", "▲ עלייה · ▼ ירידה לעומת החודש הקודם")
    change = lambda key: lambda i: ("—", None) if i == 0 else (
        f"={dref(key, FIRST + i)}-{dref(key, FIRST + i - 1)}", rows[i][key] - rows[i - 1][key])
    t1, row = table(row + 1, f"{L['revenue']} ו{L['profit']}", [
        ("חודש", "month_label", None), (L["revenue"], "revenue", MONEY), (f"שינוי ב{L['revenue']}", change("revenue"), DELTA),
        (L["profit"], "profit", MONEY), (f"שינוי ב{L['profit']}", change("profit"), DELTA), ("שיעור רווח", "margin", PCT)])
    if n > 1:
        for letter in ("D", "F"):
            cells = f"{letter}{t1 + 2}:{letter}{t1 + n}"
            ws.conditional_formatting.add(cells, CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=GOOD, bold=True)))
            ws.conditional_formatting.add(cells, CellIsRule(operator="lessThan", formula=["0"], font=Font(color=BAD, bold=True)))
    parts = [k for k in (["owner_salary"] if L["owner_salary"] else []) + [f"g{j}" for j in range(len(groups))]]
    if parts:
        names = [L["owner_salary"]] if L["owner_salary"] else []
        _, row = table(row, L["expenses"], [("חודש", "month_label", None)] +
                       [(nm, k, MONEY) for nm, k in zip(names + groups, parts)] +
                       [("שאר ההוצאות", "rest", MONEY), (f"סה״כ {L['expenses']}", "expenses", MONEY)])
    if items:
        acols = [("חודש", "month_label", None)] + [(it["label"], f"a{j}", COUNT) for j, it in enumerate(items)]
        if act["total"]:
            acols.append((act["total"], "total_activity", COUNT))
        _, row = table(row, act["title"], acols)

    # Charts, two per row
    section(row, "גרפים", "התמונה הכללית במבט אחד")
    cats_f = f"'{DATA}'!$A${FIRST}:$A${end}"
    cats = [r["label"] for r in rows]
    colnum = {k: i for i, (k, *_r) in enumerate(data_columns(data), 1)}
    charts = []

    line = _chart_frame(LineChart(), f"{L['revenue']}, {L['expenses']} ו{L['profit']} לפי חודש")
    for key, color in (("revenue", SERIES[0]), ("expenses", SERIES[1]), ("profit", SERIES[2])):
        line.series.append(_series(dsheet, colnum[key], FIRST, end, L[key], [r[key] for r in rows], cats_f, cats, color, line=True))
    charts.append((line, 21, (line.title_text, "באלפי ₪, ללא מע״מ", line_chart(d, L))))

    if n > 1:
        ch = _chart_frame(BarChart(), f"השינוי מחודש לחודש: {L['revenue']} ו{L['profit']}")
        ch.type, ch.gapWidth = "col", 120
        tf = f"'{DASH}'!$B${t1 + 2}:$B${t1 + n}"
        for j, key, color in ((3, "revenue", SERIES[0]), (5, "profit", SERIES[2])):
            vals = [rows[i][key] - rows[i - 1][key] for i in range(1, n)]
            ch.series.append(_series(ws, 1 + j, t1 + 2, t1 + n, f"שינוי ב{L[key]}", vals, tf, cats[1:], color))
        charts.append((ch, 21, (ch.title_text, "באלפי ₪", change_chart(d, L))))

    def stacked(title, keys_names):
        ch = _chart_frame(BarChart(), title)
        ch.type, ch.grouping, ch.overlap, ch.gapWidth = "col", "stacked", 100, 70
        ch.y_axis.scaling.min = 0
        for i, (key, name) in enumerate(keys_names):
            ch.series.append(_series(dsheet, colnum[key], FIRST, end, name, [row_value(r, key) for r in rows],
                                     cats_f, cats, SERIES[i % len(SERIES)]))
        return ch

    if parts:
        ch = stacked(f"{L['expenses']} לפי סוג", list(zip(parts + ["rest"], names + groups + ["שאר ההוצאות"])))
        snap = stacked_chart(d, [nm for nm, _ in expense_parts(data)], lambda r: [g(r) for _, g in expense_parts(data)], ch.title_text)
        charts.append((ch, 21, (ch.title_text, "באלפי ₪, ללא מע״מ", snap)))
    if items and act["total"]:
        ch = stacked(f"{act['title']} לפי סוג", [(f"a{j}", it["label"]) for j, it in enumerate(items)])
        snap = stacked_chart(d, [it["label"] for it in items], lambda r: r["activity"], ch.title_text, False, "")
        charts.append((ch, 21, (ch.title_text, f"המספר מעל כל עמודה: {act['total']}", snap)))
    elif items:
        for j, it in enumerate(items):
            ch = _chart_frame(BarChart(), it["label"] + (f" ({it['unit']})" if it.get("unit") else ""), height=6)
            ch.type, ch.gapWidth = "col", 80
            ch.y_axis.scaling.min = 0
            ch.legend = None
            ch.series.append(_series(dsheet, colnum[f"a{j}"], FIRST, end, it["label"], [r["activity"][j] for r in rows],
                                     cats_f, cats, SERIES[j % len(SERIES)]))
            charts.append((ch, 16, (ch.title_text, "", single_chart(d, j))))
    r = row + 1
    placements = []
    for i in range(0, len(charts), 2):
        pair = charts[i:i + 2]
        for k, (ch, _, snap) in enumerate(pair):
            ws.add_chart(ch, f"{'B' if k == 0 else 'J'}{r}")
            placements.append((1 if k == 0 else 9, r - 1, ch.width, ch.height, snap))
        r += max(h for _, h, _ in pair)
    r += 2

    # Forecast, insights, notes
    f0 = r + 1
    section(f0, "תחזית שנתית לפי הקצב הנוכחי", "הממוצע החודשי × 12. מתעדכנת עם כל חודש חדש")
    direct = f"+AVERAGE({rng('direct_costs')})" if L["direct_costs"] else ""
    by = lambda f, key: f"=INDEX({rng('month_label')},MATCH({f}({rng(key)}),{rng(key)},0))"
    fc = d["forecast"]
    facts = [("מחזור שנתי משוער", f"=AVERAGE({rng('revenue')})*12", fc["revenue"], MONEY),
             ("הוצאות שנתיות משוערות", f"=(AVERAGE({rng('expenses')}){direct})*12", fc["expenses"], MONEY),
             (f"{L['profit']} שנתי משוער", f"=AVERAGE({rng('profit')})*12", fc["profit"], MONEY),
             (f"החודש עם ה{L['revenue']} הגבוהות", by("MAX", "revenue"), d["best_revenue"], None),
             ("החודש עם הרווח הגבוה", by("MAX", "profit"), d["best_profit"], None),
             ("החודש עם הרווח הנמוך", by("MIN", "profit"), d["worst_profit"], None)]
    for i, (label, formula, value, fmt) in enumerate(facts):
        rr, c = f0 + 1 + i % 3, (2 if i < 3 else 10)
        ws.merge_cells(start_row=rr, start_column=c, end_row=rr, end_column=c + 3)
        text(f"{get_column_letter(c)}{rr}", label, 11, color=INK2)
        ws.merge_cells(start_row=rr, start_column=c + 4, end_row=rr, end_column=c + 6)
        style(cache.put(ws, f"{get_column_letter(c + 4)}{rr}", formula, value), 12, True, fmt=fmt)
        for cc in range(c, c + 7):
            ws.cell(rr, cc).border = Border(bottom=thin)

    i0 = f0 + 5
    section(i0, "תובנות לפעולה", "לשיחת הליווי")
    for i, insight in enumerate(data["insights"]):
        rr = i0 + 1 + i
        text(f"B{rr}", i + 1, 12, True, ACCENT, align="center")
        ws.merge_cells(f"C{rr}:Q{rr}")
        c = cache.put(ws, f"C{rr}", f"='{INSIGHTS}'!B{4 + i}", insight)
        c.font = Font(name=FONT, size=11, color=INK)
        c.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        ws.row_dimensions[rr].height = 30
    last = i0 + len(data["insights"])
    if data["notes"]:
        n0 = last + 2
        section(n0, "איך לקרוא את הדשבורד")
        for i, (t, x) in enumerate(data["notes"]):
            rr = n0 + 1 + i
            ws.merge_cells(f"B{rr}:D{rr}")
            text(f"B{rr}", t, 11, True, INK2)
            ws.merge_cells(f"E{rr}:Q{rr}")
            text(f"E{rr}", x, 11, color=INK)
        last = n0 + len(data["notes"])
    ws.print_area = f"A1:Q{last + 1}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return placements


def write_cached_values(path, wb, cache):
    """openpyxl writes <f> with an empty <v>. Put the calculated value in, so every viewer shows numbers."""
    index = {title: i + 1 for i, title in enumerate(wb.sheetnames)}
    by_sheet = {}
    for (title, coord), value in cache.cells.items():
        by_sheet.setdefault(index[title], {})[coord] = value
    tmp = Path(path).with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            m = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", item.filename)
            if m and int(m.group(1)) in by_sheet:
                values = by_sheet[int(m.group(1))]

                def fill_in(mo):
                    coord, attrs, formula = mo.group(1), mo.group(2), mo.group(3)
                    if coord not in values:
                        return mo.group(0)
                    v = values[coord]
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        return f'<c r="{coord}"{attrs}><f>{formula}</f><v>{repr(float(v)) if isinstance(v, float) else v}</v></c>'
                    return f'<c r="{coord}"{attrs} t="str"><f>{formula}</f><v>{xml_escape(str(v))}</v></c>'

                text = re.sub(r'<c r="([A-Z]+\d+)"((?: s="\d+")?)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>', fill_in, data.decode("utf-8"))
                data = text.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def build_xlsx(data, d, path):
    cache = Cache()
    wb = Workbook()
    wb.remove(wb.active)
    col, end = build_data_sheet(wb, data, d, cache)
    build_insights_sheet(wb, data)
    build_settings_sheet(wb, data)
    placements = build_dashboard_sheet(wb, data, d, col, end, cache)
    wb.active = 0
    wb.save(path)
    write_cached_values(path, wb, cache)
    add_chart_pictures(path, placements)


def read_xlsx(path):
    wb = load_workbook(path, data_only=False)
    config = json.loads(wb[SETTINGS]["B2"].value)
    data = {**config, "months": [], "insights": []}
    normalize({**data, "months": [{"month": "2000-01", "revenue": 0, "expenses": 0,
                                   "groups": [0] * len(data.get("expense_groups", [])),
                                   "activity": [0] * len((data.get("activity") or {}).get("items", []))}]})
    ws = wb[DATA]
    cols = data_columns(data)
    for row in range(FIRST, ws.max_row + 1):
        get = {key: ws[f"{letter}{row}"].value for key, _, _, letter in cols}
        if not get["month"]:
            continue
        zero = lambda v: num(v) if num(v) is not None else 0
        data["months"].append({
            "month": str(get["month"]), "revenue": zero(get["revenue"]), "direct_costs": zero(get.get("direct_costs")),
            "expenses": zero(get["expenses"]), "owner_salary": zero(get.get("owner_salary")),
            "groups": [zero(get[f"g{j}"]) for j in range(len(data["expense_groups"]))],
            "activity": [num(get[f"a{j}"]) for j in range(len(data["activity"]["items"]))],
            "source": get["source"] or ""})
    for r in wb[INSIGHTS].iter_rows(min_row=4, values_only=True):
        if len(r) > 1 and r[1]:
            data["insights"].append(str(r[1]))
    return data

# ---------------------------------------------------------------- HTML

def esc(s):
    return html.escape(str(s))


def money(v, signed=False):
    sign = "−" if v < 0 else ("+" if signed and v > 0 else "")
    return f'<span class="num" dir="ltr">{sign}{abs(v):,.0f}</span>&nbsp;₪'


def plain(v):
    return "—" if v is None else f"{'−' if v < 0 else ''}{abs(v):,.0f}"


def delta(v):
    if v is None:
        return '<span class="muted">—</span>'
    if abs(v) < 0.5:
        return '<span class="muted">ללא שינוי</span>'
    cls, arrow = ("up", "▲") if v > 0 else ("down", "▼")
    return f'<span class="delta {cls}">{arrow} {money(abs(v))}</span>'


def nice_ticks(lo, hi, count=5):
    lo, hi = min(lo, 0), max(hi, 0)
    if hi == lo:
        hi = lo + 1
    raw = (hi - lo) / count
    mag = 10 ** math.floor(math.log10(raw))
    step = next(m * mag for m in (1, 2, 2.5, 5, 10) if m * mag >= raw)
    start, stop = math.floor(lo / step) * step, math.ceil(hi / step) * step
    ticks, t = [], start
    while t <= stop + step / 1000:
        ticks.append(round(t, 6))
        t += step
    return ticks


def tick_label(v, thousands):
    v = v / 1000 if thousands else v
    s = f"{abs(v):,.1f}".rstrip("0").rstrip(".")
    return ("−" if v < 0 else "") + s


class Plot:
    """A plain SVG plot frame: months from right to left (RTL, like the sheet), one value axis on the right,
    hover bands with a tooltip per month."""

    def __init__(self, labels, lo, hi, thousands=True, size=(720, 280, 46, 74, 14, 34), axis=True):
        self.labels, self.thousands = labels, thousands
        self.W, self.H, axis_margin, end_margin, self.T, self.B = size
        self.L, self.R = (end_margin, axis_margin) if RTL else (axis_margin, end_margin)
        self.ticks = nice_ticks(lo, hi)
        self.lo, self.hi = self.ticks[0], self.ticks[-1]
        if not axis:
            self.ticks = [0]
        self.pw, self.ph = self.W - self.L - self.R, self.H - self.T - self.B
        self.band = self.pw / len(labels)
        self.parts = []

    def x(self, i):
        return self.W - self.R - self.band * (i + 0.5) if RTL else self.L + self.band * (i + 0.5)

    def y(self, v):
        return self.T + self.ph * (self.hi - v) / (self.hi - self.lo)

    def frame(self):
        out = []
        for t in self.ticks:
            y = self.y(t)
            out.append(f'<line x1="{self.L}" x2="{self.W - self.R}" y1="{y:.1f}" y2="{y:.1f}" class="{"base" if abs(t) < 1e-9 else "grid"}"/>')
            if len(self.ticks) > 1:
                tx, anchor = (self.W - self.R + 8, "start") if RTL else (self.L - 8, "end")
                out.append(f'<text x="{tx}" y="{y + 4:.1f}" class="tick" text-anchor="{anchor}">{tick_label(t, self.thousands)}</text>')
        for i, lab in enumerate(self.labels):
            out.append(f'<text x="{self.x(i):.1f}" y="{self.H - 10}" class="tick" text-anchor="middle">{esc(lab)}</text>')
        return out

    def bands(self, tips):
        return [f'<rect class="band" x="{self.x(i) - self.band / 2:.1f}" y="{self.T}" width="{self.band:.1f}" height="{self.ph}" '
                f'data-tip="{esc(t)}"/>' for i, t in enumerate(tips)]

    def svg(self, label):
        return (f'<svg viewBox="0 0 {self.W} {self.H}" role="img" aria-label="{esc(label)}" '
                f'preserveAspectRatio="xMidYMid meet">{"".join(self.frame() + self.parts)}</svg>')


def bar_path(x, y0, y1, w, r=4):
    """A column from baseline y0 to data end y1 with a rounded data end and a square base."""
    r = min(r, abs(y1 - y0), w / 2)
    if y1 <= y0:
        return (f"M{x:.1f},{y0:.1f}V{y1 + r:.1f}Q{x:.1f},{y1:.1f} {x + r:.1f},{y1:.1f}H{x + w - r:.1f}"
                f"Q{x + w:.1f},{y1:.1f} {x + w:.1f},{y1 + r:.1f}V{y0:.1f}Z")
    return (f"M{x:.1f},{y0:.1f}V{y1 - r:.1f}Q{x:.1f},{y1:.1f} {x + r:.1f},{y1:.1f}H{x + w - r:.1f}"
            f"Q{x + w:.1f},{y1:.1f} {x + w:.1f},{y1 - r:.1f}V{y0:.1f}Z")


def tip(label, items, unit="₪"):
    return label + "\n" + "\n".join(f"{name}: {plain(v)}{' ' + unit if unit and v is not None else ''}" for name, v in items)


def line_chart(d, L):
    rows = d["rows"]
    series = [(L["revenue"], "revenue", 1), (L["expenses"], "expenses", 2), (L["profit"], "profit", 3)]
    vals = [r[k] for _, k, _ in series for r in rows]
    p = Plot([r["label"] for r in rows], min(vals), max(vals))
    p.parts += p.bands([tip(r["label"], [(nm, r[k]) for nm, k, _ in series]) for r in rows])
    ends = []
    for _, key, slot in series:
        pts = [(p.x(i), p.y(r[key])) for i, r in enumerate(rows)]
        if len(pts) > 1:
            p.parts.append(f'<polyline class="line s{slot}" points="{" ".join(f"{x:.1f},{y:.1f}" for x, y in pts)}"/>')
        p.parts += [f'<circle class="dot s{slot}" cx="{x:.1f}" cy="{y:.1f}" r="4"/>' for x, y in pts]
        ends.append([pts[-1][1], plain(rows[-1][key])])
    ends.sort()
    for i in range(1, len(ends)):  # keep end labels apart
        ends[i][0] = max(ends[i][0], ends[i - 1][0] + 15)
    ex, anchor = (p.x(len(rows) - 1) - 10, "end") if RTL else (p.x(len(rows) - 1) + 10, "start")
    p.parts += [f'<text x="{ex:.1f}" y="{y + 4:.1f}" class="end" text-anchor="{anchor}">{t}</text>' for y, t in ends]
    return p.svg("מגמה לפי חודש"), [(nm, s) for nm, _, s in series]


def change_chart(d, L):
    rows = d["rows"][1:]
    if not rows:
        return None, []
    series = [(f"שינוי ב{L['revenue']}", "d_revenue", 1), (f"שינוי ב{L['profit']}", "d_profit", 3)]
    vals = [r[k] for _, k, _ in series for r in rows]
    p = Plot([r["label"] for r in rows], min(vals), max(vals))
    p.parts += p.bands([tip(f"{r['label']} לעומת החודש הקודם", [(nm, r[k]) for nm, k, _ in series]) for r in rows])
    w = min(24, p.band * 0.28)
    for i, r in enumerate(rows):
        for j, (_, key, slot) in enumerate(series):
            bx = p.x(i) + 1 - j * (w + 2) if RTL else p.x(i) - w - 1 + j * (w + 2)  # first series on the reading side
            p.parts.append(f'<path class="bar s{slot}" d="{bar_path(bx, p.y(0), p.y(r[key]), w)}"/>')
    return p.svg("השינוי מחודש לחודש"), [(nm, s) for nm, _, s in series]


def stacked_chart(d, names, getter, label, thousands=True, unit="₪"):
    rows = d["rows"]
    heights = [sum(v or 0 for v in getter(r)) for r in rows]
    totals = [None if None in getter(r) else h for r, h in zip(rows, heights)]  # a total only when every part is known
    p = Plot([r["label"] for r in rows], 0, max(heights) or 1, thousands)
    p.parts += p.bands([tip(r["label"], list(zip(names, getter(r))) + [("סה״כ", t)], unit) for r, t in zip(rows, totals)])
    w = min(30, p.band * 0.5)
    for i, r in enumerate(rows):
        acc, x = 0, p.x(i) - w / 2
        segs = [(v, s) for s, v in enumerate(getter(r), 1) if v]
        for k, (v, slot) in enumerate(segs):
            y0, y1 = p.y(acc), p.y(acc + v)
            gap = 0 if k == 0 else 1  # the surface gap between segments
            if k == len(segs) - 1:
                p.parts.append(f'<path class="bar s{slot}" d="{bar_path(x, y0 - gap, y1, w)}"/>')
            else:
                p.parts.append(f'<rect class="bar s{slot}" x="{x:.1f}" y="{y1 + 1:.1f}" width="{w:.1f}" height="{max(0, y0 - y1 - 1 - gap):.1f}"/>')
            acc += v
        if totals[i] is not None:
            p.parts.append(f'<text x="{p.x(i):.1f}" y="{p.y(acc) - 7:.1f}" class="tot" text-anchor="middle">'
                           f'{tick_label(acc, thousands) if thousands else f"{acc:,.0f}"}</text>')
    return p.svg(label), [(nm, s) for s, nm in enumerate(names, 1)]


def mini_charts(d, parts, is_money=True, share_of=None):
    """One small column chart per part, each on its own scale, so small items stay readable.
    parts: (name, getter, unit)."""
    rows, out = d["rows"], []
    many = len(rows) > 7
    labels = [SHORT_HE[int(r["month"][5:]) - 1] for r in rows] if many else [r["label"] for r in rows]
    for k, (name, getter, unit) in enumerate(parts):
        vals = [getter(r) for r in rows]
        known = [v for v in vals if v is not None]
        p = Plot(labels, 0, max(known or [1]) or 1, is_money, size=((700 if many else 340), 150, 6, 6, 22, 26), axis=False)
        p.parts += p.bands([f"{r['label']}\n{name}: {plain(v)}{' ' + unit if v is not None and unit else ''}"
                            for r, v in zip(rows, vals)])
        w = min(22, p.band * 0.55)
        top = max(known) if known else None
        for i, v in enumerate(vals):
            if v:
                p.parts.append(f'<path class="bar s{k % 6 + 1}" d="{bar_path(p.x(i) - w / 2, p.y(0), p.y(v), w)}"/>')
            if v is not None and (not many or i == len(vals) - 1 or v == top):
                p.parts.append(f'<text x="{p.x(i):.1f}" y="{p.y(v) - 6:.1f}" class="tot" text-anchor="middle">'
                               f'{tick_label(v, is_money) if is_money else f"{v:,.0f}"}</text>')
        avg = sum(known) / len(known) if known else None
        if avg is None:
            meta = "אין נתונים"
        elif is_money:
            meta = f"ממוצע {money(avg)}" + (f" · {sum(known) / share_of * 100:.0f}% מההוצאות" if share_of else "")
        else:
            meta = f"ממוצע {avg:,.{1 if abs(avg) < 100 else 0}f} {esc(unit or '')}".strip()
        out.append(f'<div class="mini"><div class="mini-head"><span><i class="sw s{k % 6 + 1}"></i>{esc(name)}</span>'
                   f'<span class="muted">{meta}</span></div>{p.svg(name)}</div>')
    return f'<div class="minis{" one-col" if many else ""}">' + "".join(out) + "</div>"


def legend(items):
    return '<div class="legend">' + "".join(f'<span><i class="sw s{s}"></i>{esc(nm)}</span>' for nm, s in items) + "</div>"


def chart_card(title, sub, svg_legend=None, body=None):
    if svg_legend is not None:
        svg, items = svg_legend
        if svg is None:
            return ""
        body = (legend(items) if len(items) > 1 else "") + svg
    return f'<figure class="chart"><figcaption><h3>{esc(title)}</h3><p>{esc(sub)}</p></figcaption>{body}</figure>'


def font_face():
    out = []
    for name, rng in (("heebo-hebrew.woff2", "U+0307-0308, U+0590-05FF, U+200C-2010, U+20AA, U+25CC, U+FB1D-FB4F"),
                      ("heebo-latin.woff2", "U+0000-00FF, U+0131, U+0152-0153, U+02BB-02BC, U+02C6, U+02DA, U+02DC, U+0304, "
                                            "U+0308, U+0329, U+2000-206F, U+20AC, U+2122, U+2191, U+2193, U+2212, U+2215, U+FEFF, U+FFFD")):
        f = HERE / "assets" / name
        if f.exists():
            out.append(f"@font-face{{font-family:'Heebo';font-weight:300 800;font-display:swap;"
                       f"src:url(data:font/woff2;base64,{base64.b64encode(f.read_bytes()).decode()}) format('woff2');unicode-range:{rng}}}")
    return "\n".join(out)


CSS = """
:root{--ink:#17233b;--ink2:#52514e;--muted:#7a786f;--grid:#e6e4dc;--base:#b9b7ab;--plane:#f4f3ee;--card:#ffffff;
--accent:#1f3a5f;--accent-soft:#e8eef6;--good:#0b6b2b;--good-bg:#e6f4ea;--bad:#b42318;--bad-bg:#fdecea;
--s1:#2a78d6;--s2:#eb6834;--s3:#1baf7a;--s4:#eda100;--s5:#e87ba4;--s6:#6250d6}
*{box-sizing:border-box}
html{-webkit-print-color-adjust:exact;print-color-adjust:exact}
body{margin:0;background:var(--plane);color:var(--ink);font-family:'Heebo','Segoe UI',Arial,sans-serif;font-size:15px;line-height:1.55}
.page{max-width:1000px;margin:0 auto;padding:28px 20px 48px}
.num{font-variant-numeric:tabular-nums;unicode-bidi:isolate}
.muted{color:var(--muted)}
header.hero{background:linear-gradient(135deg,#1f3a5f 0%,#2d5482 100%);color:#fff;border-radius:20px;padding:30px 34px 26px;position:relative;overflow:hidden}
header.hero::after{content:"";position:absolute;inset:auto -60px -90px auto;width:260px;height:260px;border-radius:50%;background:rgba(255,255,255,.06)}
.eyebrow{font-size:13px;letter-spacing:.04em;opacity:.8;margin:0 0 4px}
header.hero h1{font-size:34px;line-height:1.15;margin:0 0 14px;font-weight:800}
.meta{display:flex;flex-wrap:wrap;gap:8px 10px;font-size:13px}
.chip{background:rgba(255,255,255,.14);border-radius:999px;padding:4px 12px}
section{margin-top:26px}
h2{font-size:19px;margin:0 0 4px;font-weight:700;color:var(--accent)}
.sub{margin:0 0 14px;color:var(--muted);font-size:13px}
.cards{display:grid;grid-template-columns:repeat(var(--cols,4),minmax(0,1fr));gap:12px}
.card{background:var(--card);border-radius:14px;padding:16px 18px 14px;box-shadow:0 1px 2px rgba(23,35,59,.06),0 0 0 1px rgba(23,35,59,.05)}
.card .label{font-size:13px;color:var(--ink2);font-weight:500;margin:0}
.card .value{font-size:26px;font-weight:700;margin:4px 0 2px;line-height:1.2;white-space:nowrap}
.card .note{font-size:12px;color:var(--muted);margin:0}
.card.hero-card{background:var(--accent);color:#fff;box-shadow:none}
.card.hero-card .label,.card.hero-card .note{color:rgba(255,255,255,.78)}
.card.hero-card .value{font-size:30px}
.tag{display:inline-block;font-size:11px;font-weight:600;border-radius:6px;padding:1px 7px;margin-inline-start:6px;vertical-align:middle}
.tag.loss{background:#fdecea;color:var(--bad)}.tag.gain{background:var(--good-bg);color:var(--good)}
.cards.act .value{font-size:24px}
.panel+figure.chart{margin-top:14px}
.panel{background:var(--card);border-radius:16px;padding:18px 20px;box-shadow:0 1px 2px rgba(23,35,59,.06),0 0 0 1px rgba(23,35,59,.05)}
figure.chart{margin:0;background:var(--card);border-radius:16px;padding:18px 20px 10px;box-shadow:0 1px 2px rgba(23,35,59,.06),0 0 0 1px rgba(23,35,59,.05)}
figure.chart+figure.chart{margin-top:14px}
figcaption h3{margin:0;font-size:16px;font-weight:700}
figcaption p{margin:0 0 6px;font-size:12px;color:var(--muted)}
.legend{display:flex;flex-wrap:wrap;gap:4px 16px;font-size:13px;color:var(--ink2);margin:2px 0 4px}
.legend span{display:inline-flex;align-items:center;gap:6px}
.sw{display:inline-block;width:12px;height:12px;border-radius:3px}
svg{width:100%;height:auto;display:block;direction:ltr;font-family:'Heebo','Segoe UI',Arial,sans-serif}
.grid{stroke:var(--grid);stroke-width:1}.base{stroke:var(--base);stroke-width:1.2}
.tick{fill:var(--muted);font-size:12px;font-variant-numeric:tabular-nums}
.end,.tot{fill:var(--ink2);font-size:12px;font-weight:600;font-variant-numeric:tabular-nums}
.mini .tot{font-size:13px}
.line{fill:none;stroke-width:2.25;stroke-linejoin:round;stroke-linecap:round}
.dot{stroke:#fff;stroke-width:2}
.band{fill:transparent}.band:hover{fill:rgba(31,58,95,.05)}
.s1{background:var(--s1);stroke:var(--s1);fill:var(--s1)}.s2{background:var(--s2);stroke:var(--s2);fill:var(--s2)}
.s3{background:var(--s3);stroke:var(--s3);fill:var(--s3)}.s4{background:var(--s4);stroke:var(--s4);fill:var(--s4)}
.s5{background:var(--s5);stroke:var(--s5);fill:var(--s5)}.s6{background:var(--s6);stroke:var(--s6);fill:var(--s6)}
.line.s1,.line.s2,.line.s3{fill:none}.bar{stroke:none}
table{width:100%;border-collapse:collapse;font-size:14px}
th{font-size:12px;font-weight:600;color:var(--ink2);text-align:center;padding:8px 6px;border-bottom:1.5px solid var(--base);vertical-align:bottom}
td{text-align:center;padding:9px 6px;border-bottom:1px solid var(--grid);white-space:nowrap;font-variant-numeric:tabular-nums}
td:first-child,th:first-child{text-align:right;font-weight:600;color:var(--ink)}
tr:last-child td{border-bottom:0}
td.total{font-weight:700}
.delta{font-weight:600;font-size:13px;border-radius:6px;padding:2px 7px}
.delta.up{color:var(--good);background:var(--good-bg)}.delta.down{color:var(--bad);background:var(--bad-bg)}
.neg{color:var(--bad)}
.table-wrap{overflow-x:auto}
.minis{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:10px 22px;margin-top:6px}
.minis.one-col{grid-template-columns:1fr}
.mini{border-top:1px solid var(--grid);padding-top:10px}
.mini-head{display:flex;justify-content:space-between;align-items:baseline;gap:8px;font-size:13px;font-weight:600;flex-wrap:wrap}
.mini-head span:first-child{display:inline-flex;align-items:center;gap:6px}
.mini-head .muted{font-weight:400;font-size:12px}
@media screen and (max-width:760px){.minis{grid-template-columns:1fr}}
.table-wrap+.table-wrap{margin-top:18px}
.t-title{font-size:14px;font-weight:700;margin:0 0 4px;color:var(--ink)}
.forecast{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:12px}
.facts{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:12px;margin-top:12px}
.fact{background:var(--accent-soft);border-radius:12px;padding:12px 16px}
.fact p{margin:0;font-size:13px;color:var(--ink2)}.fact b{font-size:18px;color:var(--accent)}
ol.insights{list-style:none;margin:0;padding:0;counter-reset:i}
ol.insights li{counter-increment:i;position:relative;padding:12px 48px 12px 6px;border-bottom:1px solid var(--grid)}
ol.insights li:last-child{border-bottom:0}
ol.insights li::before{content:counter(i);position:absolute;right:4px;top:11px;width:28px;height:28px;border-radius:50%;
background:var(--accent);color:#fff;font-weight:700;font-size:14px;display:grid;place-items:center}
dl.notes{display:grid;grid-template-columns:max-content 1fr;gap:8px 18px;margin:0}
dl.notes dt{font-weight:700;color:var(--accent)}dl.notes dd{margin:0;color:var(--ink2)}
footer{margin-top:28px;color:var(--muted);font-size:12px;display:flex;justify-content:space-between;gap:12px;flex-wrap:wrap;
border-top:1px solid var(--grid);padding-top:12px}
#tip{position:fixed;pointer-events:none;background:#17233b;color:#fff;font-size:13px;line-height:1.5;padding:8px 12px;border-radius:8px;
white-space:pre;opacity:0;transition:opacity .12s;z-index:9;box-shadow:0 6px 18px rgba(0,0,0,.18)}
@media screen and (max-width:760px){.cards{grid-template-columns:repeat(2,minmax(0,1fr))!important}.forecast,.facts{grid-template-columns:1fr}
header.hero h1{font-size:26px}.page{padding:16px 16px 40px}.card .value{font-size:22px}}
@page{size:A4;margin:11mm 11mm 12mm}
@media print{body{background:#fff;font-size:13px}.page{max-width:none;padding:0}
.card,.panel,figure.chart{box-shadow:0 0 0 1px #e3e1d8}
section,figure.chart,.panel,.cards{break-inside:avoid}.break{break-before:page}
#tip,.screen-only{display:none}header.hero{border-radius:14px;padding:22px 26px 20px}header.hero h1{font-size:30px}
section{margin-top:18px}.card{padding:12px 14px 10px}.card .value{font-size:22px}.card.hero-card .value{font-size:24px}
figure.chart{padding:14px 16px 8px}td{padding:6px 6px}th{padding:6px}.sub{margin-bottom:10px}}
"""

JS = """
(()=>{const t=document.getElementById('tip');
document.querySelectorAll('[data-tip]').forEach(el=>{
el.addEventListener('mousemove',e=>{t.textContent=el.dataset.tip;t.style.opacity=1;
const w=t.offsetWidth,h=t.offsetHeight;let x=e.clientX+14,y=e.clientY-h-12;if(x+w>innerWidth-8)x=e.clientX-w-14;if(y<8)y=e.clientY+16;
t.style.left=x+'px';t.style.top=y+'px'});
el.addEventListener('mouseleave',()=>t.style.opacity=0)});addEventListener('scroll',()=>t.style.opacity=0,{passive:true})})();
"""


def card(label, value, note="", hero=False, tag=""):
    return (f'<div class="card{" hero-card" if hero else ""}"><p class="label">{esc(label)}{tag}</p>'
            f'<p class="value">{value}</p>{f"<p class=note>{esc(note)}</p>" if note else ""}</div>')


def grid(cards_html, count):
    return f'<div class="cards" style="--cols:{count if count <= 4 else 3}">{"".join(cards_html)}</div>'


def build_html(data, d, path):
    L, act, rows, n = data["labels"], data["activity"], d["rows"], d["n"]
    items, groups = act["items"], data["expense_groups"]
    updated = dt.date.fromisoformat(data["updated"]).strftime("%d.%m.%Y") if data.get("updated") else ""
    av, fc = d["avg"], d["forecast"]
    loss = lambda v: '<span class="tag loss">הפסד</span>' if v < 0 else ""
    per = f"ממוצע לחודש, לפי {months_count(n)}"

    money_cards = [card(L["revenue"], money(av["revenue"]), "ללא מע״מ"), card(L["expenses"], money(av["expenses"]), "ללא מע״מ"),
                   card(L["profit"], money(av["profit"]), "אחרי כל ההוצאות", hero=True, tag=loss(av["profit"]))]
    if L["owner_salary"]:
        money_cards.append(card(L["owner_salary"], money(av["owner_salary"]), "כלול בהוצאות"))
    act_cards = [card(it["label"], "—" if v is None else f'<span class="num">{v:,.{1 if abs(v) < 100 else 0}f}</span>', f"{it.get('unit') or ''} בחודש".strip())
                 for it, v in zip(items, d["avg_activity"])]

    th = lambda xs: "".join(f"<th>{esc(x)}</th>" for x in xs)
    t1 = "".join(
        f"<tr><td>{esc(r['label'])}</td><td>{money(r['revenue'])}</td><td>{delta(r['d_revenue'])}</td>"
        f"<td class=\"{'neg' if r['profit'] < 0 else ''}\">{money(r['profit'])}</td><td>{delta(r['d_profit'])}</td>"
        f"<td class=\"{'neg' if r['margin'] < 0 else ''}\"><span class=\"num\" dir=\"ltr\">{r['margin'] * 100:.1f}%</span></td></tr>"
        for r in rows)
    sec_money = (f'<section class="break"><h2>{esc(L["revenue"])} ו{esc(L["profit"])}, חודש אחרי חודש</h2>'
                 f'<p class="sub">▲ עלייה · ▼ ירידה, לעומת החודש הקודם. כל הסכומים ללא מע״מ</p>'
                 f'<div class="panel"><div class="table-wrap"><table><thead><tr>'
                 f'{th(["חודש", L["revenue"], "שינוי ב" + L["revenue"], L["profit"], "שינוי ב" + L["profit"], "שיעור רווח"])}'
                 f'</tr></thead><tbody>{t1}</tbody></table></div></div>'
                 f'{chart_card("השינוי מחודש לחודש", "באלפי ₪. מעל הקו: עלייה, מתחת לקו: ירידה", change_chart(d, L))}</section>')

    parts = expense_parts(data)
    sec_exp = ""
    if parts:
        t3 = "".join(f"<tr><td>{esc(r['label'])}</td>" + "".join(f"<td>{money(g(r))}</td>" for _, g in parts) +
                     f"<td class=\"total\">{money(r['expenses'])}</td></tr>" for r in rows)
        total_exp = sum(r["expenses"] for r in rows) or 1
        sec_exp = (f'<section class="break"><h2>{esc(L["expenses"])}</h2><p class="sub">לאן הולך הכסף בכל חודש. כל הסכומים ללא מע״מ</p>'
                   f'<div class="panel"><div class="table-wrap"><table><thead><tr>{th(["חודש"] + [p for p, _ in parts] + ["סה״כ"])}'
                   f'</tr></thead><tbody>{t3}</tbody></table></div></div>'
                   + chart_card(f"{L['expenses']} לפי סוג", "באלפי ₪, ללא מע״מ. לכל סוג גרף משלו, כדי שגם סעיפים קטנים ייראו בבירור",
                                body=mini_charts(d, [(nm, g, "₪") for nm, g in parts], True, total_exp)) + "</section>")

    sec_act = ""
    if items:
        cell = lambda v: "<td>—</td>" if v is None else f"<td>{v:,.0f}</td>"
        heads = ["חודש"] + [it["label"] for it in items] + ([act["total"]] if act["total"] else [])
        t2 = "".join(f"<tr><td>{esc(r['label'])}</td>" + "".join(cell(v) for v in r["activity"]) +
                     (f"<td class=\"total\">{plain(r['total_activity'])}</td>" if act["total"] else "") + "</tr>" for r in rows)
        if act["total"]:
            chart = chart_card(f"{act['title']} לפי סוג", f"בכל חודש. המספר מעל כל עמודה: {act['total']}",
                               stacked_chart(d, [it["label"] for it in items], lambda r: r["activity"], act["title"], False, ""))
        else:
            chart = chart_card(f"{act['title']} חודש אחרי חודש", "לכל מדד גרף משלו",
                               body=mini_charts(d, [(it["label"], (lambda j: lambda r: r["activity"][j])(j), it.get("unit") or "")
                                                    for j, it in enumerate(items)], False))
        sec_act = (f'<section class="break"><h2>{esc(act["title"])}</h2><p class="sub">חודש אחרי חודש</p>'
                   f'<div class="panel"><div class="table-wrap"><table><thead><tr>{th(heads)}</tr></thead><tbody>{t2}</tbody></table></div></div>'
                   f'{chart}</section>')

    direct_note = f"כולל {L['direct_costs']}" if L["direct_costs"] else ""
    insights = "".join(f"<li>{esc(x)}</li>" for x in data["insights"])
    notes = "".join(f"<dt>{esc(t)}</dt><dd>{esc(x)}</dd>" for t, x in data["notes"])
    doc = f"""<!doctype html>
<html lang="he" dir="rtl"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>דשבורד {esc(data['client'])} · {esc(d['period'])}</title>
<style>{font_face()}{CSS}</style></head>
<body><div class="page">
<header class="hero">
  <p class="eyebrow">דשבורד ניהולי {esc(data.get('year', ''))}</p>
  <h1>{esc(data['client'])}</h1>
  <div class="meta"><span class="chip">{esc(d['period'])} · {months_count(n)}</span><span class="chip">עודכן {updated}</span>
  <span class="chip">הוכן על ידי {esc(data.get('consultant', ''))}</span></div>
</header>
<section><h2>התמונה הכספית</h2><p class="sub">{per}</p>{grid(money_cards, len(money_cards))}</section>
{f'<section><h2>{esc(act["title"])}</h2><p class="sub">{per}</p>{grid(act_cards, len(act_cards))}</section>' if items else ''}
<section><h2>המגמה</h2><p class="sub screen-only">מעבר עם העכבר על חודש מציג את המספרים המדויקים</p>
{chart_card(f"{L['revenue']}, {L['expenses']} ו{L['profit']} לפי חודש", "באלפי ₪, ללא מע״מ", line_chart(d, L))}</section>
{sec_money}
{sec_exp}
{sec_act}
<section><h2>תחזית שנתית לפי הקצב הנוכחי</h2><p class="sub">הממוצע החודשי × 12. התחזית מתעדכנת עם כל חודש חדש</p>
<div class="forecast">{card("מחזור שנתי משוער", money(fc["revenue"]))}{card("הוצאות שנתיות משוערות", money(fc["expenses"]), direct_note)}
{card(f"{L['profit']} שנתי משוער", money(fc["profit"]), hero=True, tag=loss(fc["profit"]))}</div>
<div class="facts">
  <div class="fact"><p>החודש עם ה{esc(L['revenue'])} הגבוהות</p><b>{esc(d['best_revenue'])}</b></div>
  <div class="fact"><p>החודש עם הרווח הגבוה</p><b>{esc(d['best_profit'])}</b></div>
  <div class="fact"><p>החודש עם הרווח הנמוך</p><b>{esc(d['worst_profit'])}</b></div>
</div></section>
{f'<section><h2>תובנות לפעולה</h2><p class="sub">לשיחת הליווי</p><div class="panel"><ol class="insights">{insights}</ol></div></section>' if insights else ''}
{f'<section><h2>איך לקרוא את הדשבורד</h2><div class="panel"><dl class="notes">{notes}</dl></div></section>' if notes else ''}
<footer><span>{esc(data['client'])} · {esc(d['period'])}</span><span>הוכן על ידי {esc(data.get('consultant', ''))} · מקור: הניתוחים הפיננסיים החודשיים</span></footer>
</div><div id="tip" role="tooltip"></div><script>{JS}</script></body></html>"""
    Path(path).write_text(doc, encoding="utf-8")


def single_chart(d, j):
    """One activity measure as columns, for the Excel picture of a small chart."""
    rows = d["rows"]
    vals = [r["activity"][j] for r in rows]
    known = [v for v in vals if v is not None]
    p = Plot([r["label"] for r in rows], 0, max(known or [1]) or 1, False, size=(720, 230, 46, 20, 14, 30))
    w = min(30, p.band * 0.5)
    for i, v in enumerate(vals):
        if v:
            p.parts.append(f'<path class="bar s{j % 6 + 1}" d="{bar_path(p.x(i) - w / 2, p.y(0), p.y(v), w)}"/>')
            p.parts.append(f'<text x="{p.x(i):.1f}" y="{p.y(v) - 6:.1f}" class="tot" text-anchor="middle">{v:,.0f}</text>')
    return p.svg("פעילות"), []


SNAP_CSS = """body{margin:0;background:#fff;overflow:hidden}.snap{box-sizing:border-box;padding:12px 16px 6px;display:flex;flex-direction:column}
.snap h3{margin:0;font-size:15px;font-weight:700;color:#17233b}.snap p{margin:0 0 4px;font-size:11px;color:#7a786f}
.snap .legend{margin:0 0 2px}.snap .plot{flex:1;min-height:0;display:flex;justify-content:center}.snap svg{height:100%;width:auto;max-width:100%}"""


def render_pictures(placements):
    """PNG of every chart, drawn like the report, so viewers that skip Excel charts still show them. None without a browser."""
    exe = find_browser()
    if not exe:
        return None
    out = []
    with tempfile.TemporaryDirectory() as tmp:
        for i, (_, _, w_cm, h_cm, (title, sub, (svg, items))) in enumerate(placements):
            w, h = round(w_cm / 2.54 * 96), round(h_cm / 2.54 * 96)
            page = Path(tmp) / f"c{i}.html"
            page.write_text(
                f'<!doctype html><html lang="he" dir="rtl"><head><meta charset="utf-8"><style>{font_face()}{CSS}{SNAP_CSS}'
                f'.snap{{width:{w}px;height:{h}px}}</style></head><body><div class="snap"><h3>{esc(title)}</h3>'
                f'{f"<p>{esc(sub)}</p>" if sub else ""}{legend(items) if len(items) > 1 else ""}<div class="plot">{svg}</div></div></body></html>',
                encoding="utf-8")
            png = Path(tmp) / f"c{i}.png"
            proc = subprocess.Popen([exe, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--no-first-run",
                                     "--no-default-browser-check", f"--user-data-dir={Path(tmp) / 'profile'}",
                                     f"--window-size={w},{h}", "--force-device-scale-factor=2", f"--screenshot={png}",
                                     page.as_uri()], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            try:
                proc.wait(timeout=60)
            except subprocess.TimeoutExpired:
                proc.kill()
            if not png.exists():
                return None
            out.append(png.read_bytes())
    return out


def add_chart_pictures(path, placements):
    """Put a picture of each chart under it (charts are drawn on top in Excel and Numbers; previews show the picture)."""
    if not placements:
        return
    pngs = render_pictures(placements)
    if not pngs:
        return
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    anchors = []
    for i, ((col, row, w_cm, h_cm, (title, *_)), _) in enumerate(zip(placements, pngs)):
        # a little smaller than the chart and centred under it, so the chart always covers it
        cx, cy = round(w_cm * 360000 * 0.86), round(h_cm * 360000 * 0.86)
        dx, dy = round(w_cm * 360000 * 0.07), round(h_cm * 360000 * 0.07)
        anchors.append(
            f'<oneCellAnchor><from><col>{col}</col><colOff>{dx}</colOff><row>{row}</row><rowOff>{dy}</rowOff></from>'
            f'<ext cx="{cx}" cy="{cy}"/><pic><nvPicPr><cNvPr id="{1000 + i}" name="תמונת גרף {i + 1}" descr="{xml_escape(title)}"/>'
            f'<cNvPicPr><a:picLocks xmlns:a="{A}" noChangeAspect="1"/></cNvPicPr></nvPicPr>'
            f'<blipFill><a:blip xmlns:a="{A}" xmlns:r="{R}" r:embed="rIdPic{i + 1}"/><a:stretch xmlns:a="{A}"><a:fillRect/></a:stretch></blipFill>'
            f'<spPr><a:xfrm xmlns:a="{A}"><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
            f'<a:prstGeom xmlns:a="{A}" prst="rect"><a:avLst/></a:prstGeom></spPr></pic><clientData/></oneCellAnchor>')
    rels = "".join(f'<Relationship Type="{R}/image" Target="/xl/media/dashboard-chart{i + 1}.png" Id="rIdPic{i + 1}"/>'
                   for i in range(len(pngs)))
    tmp = Path(path).with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename).decode("utf-8") if item.filename.endswith((".xml", ".rels")) else zin.read(item.filename)
            if item.filename == "xl/drawings/drawing1.xml":  # the dashboard is the only sheet with charts
                start = data.index(">", data.index("<wsDr")) + 1
                data = data[:start] + "".join(anchors) + data[start:]  # pictures first, so charts are drawn over them
            elif item.filename == "xl/drawings/_rels/drawing1.xml.rels":
                data = data.replace("</Relationships>", rels + "</Relationships>")
            elif item.filename == "[Content_Types].xml" and 'Extension="png"' not in data:
                data = data.replace("<Default ", '<Default Extension="png" ContentType="image/png"/><Default ', 1)
            zout.writestr(item, data.encode("utf-8") if isinstance(data, str) else data)
        for i, png in enumerate(pngs):
            zout.writestr(f"xl/media/dashboard-chart{i + 1}.png", png)
    tmp.replace(path)


def find_browser():
    candidates = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
        "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
        "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    ]
    for c in candidates:
        if Path(c).exists():
            return c
    for name in ("msedge", "google-chrome", "chromium", "chromium-browser", "chrome"):
        if shutil.which(name):
            return shutil.which(name)
    return None


def build_pdf(html_path, pdf_path):
    exe = find_browser()
    if not exe:
        sys.exit("לא נמצא Edge או Chrome ליצירת PDF. אפשר לפתוח את קובץ ה-HTML ולהדפיס ל-PDF (Ctrl+P).")
    pdf_path = Path(pdf_path).resolve()
    with tempfile.TemporaryDirectory() as profile:  # own profile: works while the user's browser is open
        subprocess.run([exe, "--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--print-to-pdf-no-header",
                        f"--user-data-dir={profile}", f"--print-to-pdf={pdf_path}", Path(html_path).resolve().as_uri()],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)
    if not pdf_path.exists():
        sys.exit("יצירת ה-PDF נכשלה")

# ---------------------------------------------------------------- CLI

def main():
    for stream in (sys.stdout, sys.stderr):  # a Windows console in cp1252/cp1255 cannot print ✓ or Hebrew
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build", help="create or update the Excel dashboard from a data file")
    b.add_argument("data")
    b.add_argument("--xlsx", required=True)
    r = sub.add_parser("read", help="the Excel dashboard back as a data file")
    r.add_argument("xlsx")
    r.add_argument("--json")
    p = sub.add_parser("report", help="the presentation report (PDF and/or HTML) from the Excel as it is now")
    p.add_argument("xlsx")
    p.add_argument("--pdf")
    p.add_argument("--html")
    a = ap.parse_args()
    if a.cmd == "read":
        out = json.dumps(read_xlsx(a.xlsx), ensure_ascii=False, indent=2)
        Path(a.json).write_text(out, encoding="utf-8") if a.json else print(out)
        return
    if a.cmd == "build":
        data = normalize(json.loads(Path(a.data).read_text(encoding="utf-8")))
        d = derive(data)
        build_xlsx(data, d, a.xlsx)
    else:
        if not (a.pdf or a.html):
            sys.exit("צריך --pdf או --html")
        data = normalize(read_xlsx(a.xlsx))
        d = derive(data)
        if a.html:
            build_html(data, d, a.html)
        if a.pdf:
            with tempfile.TemporaryDirectory() as tmp:
                html_path = a.html or str(Path(tmp) / "report.html")
                if not a.html:
                    build_html(data, d, html_path)
                build_pdf(html_path, a.pdf)
    print(f"✓ {data['client']}: {d['period']} ({months_count(d['n'])})")


if __name__ == "__main__":
    main()
